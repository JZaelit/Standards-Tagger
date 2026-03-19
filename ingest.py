"""
standards_ingest/ingest.py
--------------------------
Word document (.docx) ingestion pipeline for standards documents.
Replaces the PDF pipeline entirely — no spatial reconstruction,
no zone cropping, no vision fallback needed.

Pipeline per document:
  1. Walk paragraphs in document order
  2. Detect standard ID at start of each paragraph
  3. Accumulate description text until the next standard ID
  4. Extract table-based standards (NGSS style)
  5. Output clean JSON ready for structure.py

Usage:
    python ingest.py                             # process all docs in CORPUS
    python ingest.py --doc docs/ELA_CCSS.docx --subject ELA
    python ingest.py --dry-run                   # parse + print, no file output
    python ingest.py --doc docs/ELA_CCSS.docx --subject ELA --dry-run
"""

import argparse
import json
import re
from pathlib import Path

from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph
from docx.oxml.ns import qn


# ── CONFIGURATION ──────────────────────────────────────────────────────────────

CORPUS = [
    {"path": "docs/ELA_CCSS_Standards1.docx",  "subject": "ELA"},
    {"path": "docs/ELA-CA.docx",                "subject": "ELA"},
    {"path": "docs/MATH_CCSS_Standards1.docx",  "subject": "Math"},
    {"path": "docs/MATH-CA.docx",               "subject": "Math"},
    {"path": "docs/HISTORY-CA.docx",            "subject": "History-SS"},
    {"path": "docs/NGSS Combined.docx",         "subject": "Science"},
    {"path": "docs/SCIENCEearth-CA.docx",       "subject": "Science"},
    {"path": "docs/SCIENCElife-CA.docx",        "subject": "Science"},
    {"path": "docs/SCIENCEphysical-CA.docx",    "subject": "Science"},
]

OUTPUT_DIR = Path("output")

# Standard code patterns per subject
STANDARD_PATTERNS = {
    "ELA": re.compile(
        r'^(RL|RI|RF|W|SL|L|RH|RST|WHST)\.'
        r'(K|[1-9]|1[0-2])\.'
        r'(\d+[a-z]?)\b'
    ),
    "Math": re.compile(
        r'^([A-Z]{1,6}(?:\.[A-Z]+)?)\.'
        r'([A-Z]\.)?'
        r'(\d+)\b'
    ),
    "Science": re.compile(
        r'^(\d+|[A-Z]{2,4})\s*[-\u2013]\s*'
        r'([A-Z]{2,4}S)\s*[-\u2013]\s*'
        r'(\d+)\b'
    ),
    "History-SS": re.compile(
        r'^(\d+)\s*\.\s*(\d+(?:\.\d+)?)\b'
    ),
}

GRADE_BAND_MAP = {
    "K": "K-2",  "1": "K-2",  "2": "K-2",
    "3": "3-5",  "4": "3-5",  "5": "3-5",
    "6": "6-8",  "7": "6-8",  "8": "6-8",
    "9": "9-10", "10": "9-10",
    "11": "11-12", "12": "11-12",
}


# ── HELPERS ────────────────────────────────────────────────────────────────────

def clean_text(text: str) -> str:
    text = text.replace("\xa0", " ")
    text = text.replace("\u2013", "-")
    text = text.replace("\u2014", "-")
    text = re.sub(r"\s{2,}", " ", text)
    return text.strip()


def is_heading_or_noise(text: str) -> bool:
    noise_patterns = [
        r"^(reading|writing|speaking|listening|language|foundational)",
        r"^(key ideas|craft and structure|integration|range of|text types)",
        r"^(standards for|common core|california|grade \d|grades \d)",
        r"^(anchor standard|college and career)",
        r"^\d+$",
        r"^[ivxlcdm]+$",
    ]
    lower = text.lower()
    return any(re.match(p, lower) for p in noise_patterns)


def get_style_name(para: Paragraph) -> str:
    try:
        return para.style.name if para.style else ""
    except Exception:
        return ""


def iter_block_items(doc: Document):
    """
    Yield paragraphs and tables in document order.
    doc.paragraphs alone skips tables — this doesn't.
    """
    body = doc.element.body
    for child in body:
        if child.tag == qn("w:p"):
            yield Paragraph(child, doc)
        elif child.tag == qn("w:tbl"):
            yield Table(child, doc)


# ── STANDARD ID DETECTION ──────────────────────────────────────────────────────

def detect_standard_id(text: str, subject: str):
    pattern = STANDARD_PATTERNS.get(subject)
    if not pattern:
        return None
    return pattern.match(text)


def parse_standard_id(match, subject: str) -> dict:
    groups = match.groups()

    if subject == "ELA":
        strand, grade, number = groups[0], groups[1], groups[2]
        standard_id = f"{strand}.{grade}.{number}"

    elif subject == "Math":
        domain = groups[0]
        cluster = groups[1].rstrip(".") if groups[1] else None
        number = groups[2]
        standard_id = f"{domain}.{cluster}.{number}" if cluster else f"{domain}.{number}"
        strand, grade = domain, ""

    elif subject == "Science":
        grade_part, category, number = groups[0], groups[1], groups[2]
        standard_id = f"{grade_part}-{category}-{number}"
        strand, grade = category, grade_part

    elif subject == "History-SS":
        major, minor = groups[0], groups[1]
        standard_id = f"{major}.{minor}"
        strand, grade = major, major

    else:
        standard_id = match.group(0)
        strand, grade, number = "", "", ""

    return {
        "standard_id":     standard_id,
        "strand":          strand,
        "grade":           str(grade),
        "standard_number": str(groups[-1]) if groups else "",
        "grade_band":      GRADE_BAND_MAP.get(str(grade), ""),
    }


# ── PARAGRAPH EXTRACTION ───────────────────────────────────────────────────────

def extract_from_paragraphs(doc: Document, subject: str) -> list[dict]:
    standards = []
    current = None

    for item in iter_block_items(doc):

        # Tables
        if isinstance(item, Table):
            table_standards = extract_from_table(item, subject)
            if table_standards:
                if current:
                    standards.append(finalize_standard(current))
                    current = None
                standards.extend(table_standards)
            continue

        # Paragraphs
        para = item
        raw_text = para.text
        if not raw_text:
            continue

        text = clean_text(raw_text)
        if not text:
            continue

        style = get_style_name(para)

        # Skip heading paragraphs that aren't standards
        if "Heading" in style and not detect_standard_id(text, subject):
            if current:
                standards.append(finalize_standard(current))
                current = None
            continue

        match = detect_standard_id(text, subject)

        if match:
            # Save previous standard
            if current:
                standards.append(finalize_standard(current))

            fields = parse_standard_id(match, subject)
            description = text[match.end():].strip(" .")

            current = {
                **fields,
                "description":        description,
                "continuation_lines": [],
                "na_flag":            "not applicable" in text.lower(),
                "is_sub_item":        bool(re.search(r'\.\d+[a-z]$', fields["standard_id"])),
            }

        elif current and text and not is_heading_or_noise(text):
            # Continuation text — append to current standard's description
            current["continuation_lines"].append(text)

    if current:
        standards.append(finalize_standard(current))

    return standards


def finalize_standard(current: dict) -> dict:
    description = current["description"]

    for line in current.get("continuation_lines", []):
        if is_heading_or_noise(line):
            break
        description += " " + line

    description = re.sub(r"\s{2,}", " ", description).strip()

    return {
        "standard_id":     current["standard_id"],
        "strand":          current.get("strand", ""),
        "grade":           current.get("grade", ""),
        "grade_band":      current.get("grade_band", ""),
        "standard_number": current.get("standard_number", ""),
        "description":     description,
        "na_flag":         current.get("na_flag", False),
        "is_sub_item":     current.get("is_sub_item", False),
    }


# ── TABLE EXTRACTION ───────────────────────────────────────────────────────────

def extract_from_table(table: Table, subject: str) -> list[dict]:
    """
    Scan every cell in a table for standard IDs.
    Each matching cell becomes its own standard entry.
    Used for NGSS performance expectation tables.
    """
    standards = []

    for row in table.rows:
        for cell in row.cells:
            cell_text = clean_text(cell.text)
            if not cell_text:
                continue

            match = detect_standard_id(cell_text, subject)
            if not match:
                continue

            fields = parse_standard_id(match, subject)
            description = cell_text[match.end():].strip(" .")

            # Pull additional paragraphs from the cell
            extra = [
                clean_text(p.text)
                for p in cell.paragraphs
                if clean_text(p.text) and clean_text(p.text) != cell_text
            ]
            if extra:
                description += " " + " ".join(extra)

            standards.append({
                **fields,
                "description": re.sub(r"\s{2,}", " ", description).strip(),
                "na_flag":     "not applicable" in cell_text.lower(),
                "is_sub_item": False,
                "from_table":  True,
            })

    return standards


# ── DEDUPLICATION ──────────────────────────────────────────────────────────────

def dedupe(standards: list[dict]) -> list[dict]:
    seen: dict[str, dict] = {}
    for std in standards:
        sid = std["standard_id"]
        if sid not in seen:
            seen[sid] = std
        elif len(std.get("description", "")) > len(seen[sid].get("description", "")):
            seen[sid] = std
    return list(seen.values())


# ── MAIN EXTRACTION ────────────────────────────────────────────────────────────

def ingest_docx(docx_path: str, subject: str, verbose: bool = True) -> dict:
    path = Path(docx_path)

    if not path.exists():
        return {"error": f"File not found: {docx_path}", "standards": []}

    if verbose:
        print(f"\n{'='*55}")
        print(f"  {path.name}  (subject={subject})")
        print(f"{'='*55}")

    doc = Document(str(path))
    total_paras  = len(doc.paragraphs)
    total_tables = len(doc.tables)

    if verbose:
        print(f"  Paragraphs: {total_paras}   Tables: {total_tables}")

    raw_standards = extract_from_paragraphs(doc, subject)
    standards     = dedupe(raw_standards)

    for std in standards:
        std["subject"]         = subject
        std["source_doc"]      = path.stem
        std["is_ca_extension"] = "-CA" in path.stem or path.stem.endswith("CA")

    unique_ids = sorted(set(s["standard_id"] for s in standards))

    if verbose:
        print(f"  Raw standards:  {len(raw_standards)}")
        print(f"  After dedupe:   {len(standards)}")
        print(f"  Unique IDs:     {len(unique_ids)}")
        if unique_ids:
            sample = unique_ids[:8]
            ellipsis = " ..." if len(unique_ids) > 8 else ""
            print(f"  Sample IDs:     {', '.join(sample)}{ellipsis}")
        if len(standards) < 20:
            print(f"  ⚠  Very few standards — check STANDARD_PATTERNS for {subject}")

    # Output shape is identical to the old PDF ingest.py
    # so structure.py works without any changes
    return {
        "source":                   str(docx_path),
        "subject":                  subject,
        "total_paragraphs":         total_paras,
        "total_tables":             total_tables,
        "total_standards_found":    len(standards),
        "total_standard_ids_found": len(unique_ids),
        "standard_ids":             unique_ids,
        "pages": [{
            "page_number": 1,
            "source":      str(docx_path),
            "subject":     subject,
            "strategy":    "docx",
            "word_count":  len(standards),
            "raw_text":    "\n\n".join(
                f"{s['standard_id']} {s['description']}" for s in standards
            ),
            "standards":   standards,
        }],
    }


# ── DRY RUN ────────────────────────────────────────────────────────────────────

def dry_run_doc(docx_path: str, subject: str):
    path = Path(docx_path)
    if not path.exists():
        print(f"  NOT FOUND: {docx_path}")
        return

    result = ingest_docx(docx_path, subject, verbose=True)

    if "error" in result:
        print(f"  Error: {result['error']}")
        return

    standards = result["pages"][0].get("standards", [])
    print(f"\n  First 10 parsed standards:")
    for std in standards[:10]:
        flags = ""
        if std.get("na_flag"):     flags += " [N/A]"
        if std.get("is_sub_item"): flags += " [sub-item]"
        desc_preview = std.get("description", "")[:70]
        ellipsis = "..." if len(std.get("description", "")) > 70 else ""
        print(f"    {std['standard_id']:<15}{flags}")
        print(f"    → {desc_preview}{ellipsis}\n")


# ── CLI ─────────────────────────────────────────────────────────────────────────

def main():
    parser = argparse.ArgumentParser(description="Word doc standards ingestion pipeline")
    parser.add_argument("--doc",     help="Path to a single .docx file")
    parser.add_argument("--subject", help="Subject (ELA, Math, Science, History-SS)")
    parser.add_argument("--output",  default="output", help="Output directory for JSON")
    parser.add_argument("--dry-run", action="store_true", help="Parse and print — no file output")
    args = parser.parse_args()

    output_dir = Path(args.output)
    output_dir.mkdir(exist_ok=True)

    corpus = CORPUS
    if args.doc:
        corpus = [{"path": args.doc, "subject": args.subject or "ELA"}]

    if args.dry_run:
        print("\n🔍 DRY RUN\n")
        for entry in corpus:
            dry_run_doc(entry["path"], entry["subject"])
        return

    all_results = []

    for entry in corpus:
        result = ingest_docx(entry["path"], entry["subject"])

        if "error" in result:
            print(f"  ✗ {result['error']}")
            continue

        all_results.append(result)
        stem     = Path(entry["path"]).stem
        out_path = output_dir / f"{stem}.json"
        with open(out_path, "w", encoding="utf-8") as f:
            json.dump(result, f, indent=2, ensure_ascii=False)
        print(f"  → Saved: {out_path}")

    if len(all_results) > 1:
        combined = output_dir / "all_standards_raw.json"
        with open(combined, "w", encoding="utf-8") as f:
            json.dump(all_results, f, indent=2, ensure_ascii=False)
        total = sum(r.get("total_standard_ids_found", 0) for r in all_results)
        print(f"\n✅ Combined output: {combined}")
        print(f"   Total standard IDs: {total}")


if __name__ == "__main__":
    main()